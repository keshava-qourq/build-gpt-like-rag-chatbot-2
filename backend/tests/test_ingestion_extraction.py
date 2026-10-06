"""Per-format extraction fidelity/provenance and crash-recovery/concurrency
config for the ingestion worker (AC-034..038, AC-047..049)."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

import pytest
from celery.exceptions import Retry

from app import ingestion, queue
from app.database import SessionLocal
from app.models import Chunk, Document, Organization, User


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def org(db):
    organization = Organization(id=uuid.uuid4(), name="Test Org", created_at=datetime.now(UTC))
    db.add(organization)
    db.commit()
    return organization


@pytest.fixture
def user(db, org):
    u = User(
        id=uuid.uuid4(),
        org_id=org.id,
        email=f"{uuid.uuid4()}@example.com",
        password_hash="x",
        role="member",
        is_active=True,
        created_at=datetime.now(UTC),
    )
    db.add(u)
    db.commit()
    return u


def _make_document(db, org, user, *, fmt="txt", s3_key="any-key", status="queued"):
    document = Document(
        id=uuid.uuid4(),
        org_id=org.id,
        uploader_id=user.id,
        filename="file." + fmt,
        format=fmt,
        size_bytes=10,
        s3_key=s3_key,
        status=status,
        created_at=datetime.now(UTC),
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


class _FakeEmbeddingsProvider:
    model_name = "fake-model"
    dimension = 3

    async def embed(self, texts):
        return [[0.1, 0.2, 0.3] for _ in texts]


def _build_pdf_bytes(pages: list[str]) -> bytes:
    import fitz

    doc = fitz.open()
    for page_text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), page_text)
    data = doc.tobytes()
    doc.close()
    return data


def _build_docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("Intro paragraph before the table.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Row1 A"
    table.cell(1, 1).text = "Row1 B"
    document.add_paragraph("Outro paragraph after the table.")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# AC-034: PDF page-by-page extraction with page provenance.
# ---------------------------------------------------------------------------


def test_pdf_extraction_captures_every_page_with_provenance():
    pdf_bytes = _build_pdf_bytes(["Page one content", "Page two content", "Page three content"])
    passages = ingestion.extract_passages("pdf", pdf_bytes)

    assert [p.page for p in passages] == [1, 2, 3]
    assert "Page one" in passages[0].text
    assert "Page two" in passages[1].text
    assert "Page three" in passages[2].text


def test_pdf_chunks_record_their_page_start_and_end(db, org, user, monkeypatch):
    pdf_bytes = _build_pdf_bytes(["Alpha page text", "Beta page text"])
    document = _make_document(db, org, user, fmt="pdf")

    monkeypatch.setattr(ingestion, "_fetch_object", lambda key: pdf_bytes)
    monkeypatch.setattr(ingestion, "get_embeddings_provider", lambda: _FakeEmbeddingsProvider())

    ingestion.ingest_document(str(document.id))

    db.expire_all()
    chunks = (
        db.query(Chunk)
        .filter(Chunk.document_id == document.id)
        .order_by(Chunk.ordinal)
        .all()
    )
    assert len(chunks) == 2
    assert chunks[0].page_start == 1
    assert chunks[0].page_end == 1
    assert chunks[1].page_start == 2
    assert chunks[1].page_end == 2


# ---------------------------------------------------------------------------
# AC-035: DOCX paragraphs and tables in document order.
# ---------------------------------------------------------------------------


def test_docx_extraction_includes_paragraphs_and_table_in_order():
    docx_bytes = _build_docx_bytes()
    passages = ingestion.extract_passages("docx", docx_bytes)
    assert len(passages) == 1
    text = passages[0].text

    intro_idx = text.index("Intro paragraph")
    header_idx = text.index("Header A")
    row_idx = text.index("Row1 A")
    outro_idx = text.index("Outro paragraph")

    assert intro_idx < header_idx < row_idx < outro_idx
    assert "Header B" in text
    assert "Row1 B" in text


# ---------------------------------------------------------------------------
# AC-036: CSV header + rows with row-range provenance.
# ---------------------------------------------------------------------------


def test_csv_extraction_yields_header_and_rows_with_row_numbers():
    csv_bytes = b"name,age\nAlice,30\nBob,40\n"
    passages = ingestion.extract_passages("csv", csv_bytes)

    assert [p.row_start for p in passages] == [1, 2, 3]
    assert [p.row_end for p in passages] == [1, 2, 3]
    assert passages[0].text == "name, age"
    assert passages[1].text == "Alice, 30"


def test_csv_chunks_record_their_row_range(db, org, user, monkeypatch):
    csv_bytes = b"name,age\nAlice,30\nBob,40\n"
    document = _make_document(db, org, user, fmt="csv")

    monkeypatch.setattr(ingestion, "_fetch_object", lambda key: csv_bytes)
    monkeypatch.setattr(ingestion, "get_embeddings_provider", lambda: _FakeEmbeddingsProvider())

    ingestion.ingest_document(str(document.id))

    db.expire_all()
    chunks = (
        db.query(Chunk)
        .filter(Chunk.document_id == document.id)
        .order_by(Chunk.ordinal)
        .all()
    )
    assert len(chunks) == 3
    assert [(c.row_start, c.row_end) for c in chunks] == [(1, 1), (2, 2), (3, 3)]
    assert chunks[0].page_start is None


# ---------------------------------------------------------------------------
# AC-037: TXT / Markdown verbatim, including headings and fenced code.
# ---------------------------------------------------------------------------


def test_markdown_structure_preserved_verbatim():
    md = b"# Heading\n\nSome text.\n\n```python\nprint('hi')\n```\n"
    passages = ingestion.extract_passages("md", md)
    assert len(passages) == 1
    assert passages[0].text == md.decode("utf-8")
    assert "# Heading" in passages[0].text
    assert "```python" in passages[0].text


def test_txt_extraction_verbatim():
    text = b"plain text\nwith two lines"
    passages = ingestion.extract_passages("txt", text)
    assert passages[0].text == text.decode("utf-8")


# ---------------------------------------------------------------------------
# AC-038: non-English content, no special-casing, no extra failure path.
# ---------------------------------------------------------------------------


def test_non_english_document_extracts_chunks_and_indexes(db, org, user, monkeypatch):
    japanese_text = "これはテスト文書です。" * 5
    document = _make_document(db, org, user, fmt="txt")

    monkeypatch.setattr(
        ingestion, "_fetch_object", lambda key: japanese_text.encode("utf-8")
    )
    monkeypatch.setattr(ingestion, "get_embeddings_provider", lambda: _FakeEmbeddingsProvider())

    ingestion.ingest_document(str(document.id))

    db.expire_all()
    refreshed = db.query(Document).filter(Document.id == document.id).one()
    assert refreshed.status == "ready"
    chunks = db.query(Chunk).filter(Chunk.document_id == document.id).all()
    assert len(chunks) >= 1
    assert japanese_text[:10] in chunks[0].text


# ---------------------------------------------------------------------------
# AC-048: worker concurrency read from configuration.
# ---------------------------------------------------------------------------


def test_celery_worker_concurrency_is_configurable(monkeypatch):
    monkeypatch.setenv("CELERY_WORKER_CONCURRENCY", "7")
    import importlib

    reloaded = importlib.reload(queue)
    try:
        assert reloaded.CELERY_WORKER_CONCURRENCY == 7
        assert reloaded.celery_app.conf.worker_concurrency == 7
    finally:
        importlib.reload(queue)


def test_celery_task_uses_acks_late_and_rejects_on_worker_lost():
    assert queue.celery_app.conf.task_acks_late is True
    assert queue.celery_app.conf.task_reject_on_worker_lost is True


# ---------------------------------------------------------------------------
# AC-049: bounded retry on crash recovery, never stuck in 'processing'.
# ---------------------------------------------------------------------------


def test_redelivered_task_under_max_retries_is_retried(db, org, user):
    document = _make_document(db, org, user, fmt="txt", status="processing")

    # `.apply(..., retries=N)` is Celery's own eager-execution hook for
    # simulating a task that has already been redelivered N times -- a
    # direct call to the task always runs with a fresh retries=0 request,
    # which is not what a redelivered task looks like.
    with pytest.raises(Retry):
        ingestion.ingest_document.apply(args=[str(document.id)], retries=0, throw=True)

    db.expire_all()
    refreshed = db.query(Document).filter(Document.id == document.id).one()
    assert refreshed.status == "processing"


def test_redelivered_task_exceeding_max_retries_marks_failed(db, org, user):
    document = _make_document(db, org, user, fmt="txt", status="processing")

    ingestion.ingest_document.apply(
        args=[str(document.id)], retries=ingestion.MAX_INGEST_RETRIES
    )

    db.expire_all()
    refreshed = db.query(Document).filter(Document.id == document.id).one()
    assert refreshed.status == "failed"
    assert refreshed.failure_reason
    assert "retries" in refreshed.failure_reason or "crash" in refreshed.failure_reason
